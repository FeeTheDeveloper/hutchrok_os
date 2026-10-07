"""Build the repo-owned Texas filing correction preparation workbook."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "filing-department" / "Texas Filing Correction Preparation Workbook.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def mark_row_as_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    mark_row_as_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_shading(cell, "1F4E78")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
            run.font.size = Pt(9)
        if widths:
            cell.width = Inches(widths[index])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.text = value
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=130, bottom=130)
            if row_index % 2:
                set_cell_shading(cell, "F4F7FA")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
            if widths:
                cell.width = Inches(widths[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_checkbox_line(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run("[ ] ").bold = True
    paragraph.add_run(text)


def add_fill_line(doc, label, value="UNRESOLVED"):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(f"{label}: ")
    run.bold = True
    paragraph.add_run(value)


def remove_paragraph_borders(paragraph_or_style):
    element = paragraph_or_style._element
    p_pr = element.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(24)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
remove_paragraph_borders(styles["Title"])
for style_name, size in (("Heading 1", 16), ("Heading 2", 12)):
    styles[style_name].font.name = "Aptos Display"
    styles[style_name].font.size = Pt(size)
    styles[style_name].font.bold = True
    styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
    styles[style_name].paragraph_format.space_before = Pt(10)
    styles[style_name].paragraph_format.space_after = Pt(5)

title = doc.add_paragraph(style="Title")
title.add_run("Texas Filing Correction Preparation Workbook")
remove_paragraph_borders(title)
subtitle = doc.add_paragraph()
subtitle.add_run("Secretary of State and Comptroller review packet").bold = True
subtitle.paragraph_format.space_after = Pt(12)

intro = doc.add_paragraph()
intro.add_run("Purpose. ").bold = True
intro.add_run(
    "Use this workbook to audit a Texas entity record, document address or filing corrections, "
    "and prepare an unsigned packet for authorized human review. It does not authorize or perform "
    "login, signature, attestation, checkout, payment, transmission, or submission."
)

doc.add_heading("Case control", level=1)
add_table(
    doc,
    ["Field", "Controlled value"],
    [
        ("Opaque case ID", "UNASSIGNED"),
        ("Tenant and environment", "UNASSIGNED / local"),
        ("Case status", "BLOCKED_MISSING_SOURCE"),
        ("Packet version and SHA-256", "v0 / NOT CALCULATED"),
        ("Prepared by and date", "UNASSIGNED"),
        ("Reviewed by and date", "UNASSIGNED"),
    ],
    [2.2, 5.0],
)

doc.add_heading("Required source facts", level=1)
for item in (
    "Exact legal entity name and entity type",
    "Texas SOS file number and original filing date",
    "Current agency record or immutable portal preview",
    "Authorized requester and basis of authority",
    "Exact current address fields and exact proposed address fields",
    "Evidence supporting each proposed value",
    "Authorized signer role without applying a signature",
):
    add_checkbox_line(doc, item)

warning = doc.add_paragraph()
warning.add_run("Restricted data. ").bold = True
warning.add_run(
    "Keep credentials, MFA codes, full portal record URLs, payment details, taxpayer identifiers, "
    "veteran verification codes, DD-214 records, signatures, and unnecessary personal data out of Git."
)

doc.add_section(WD_SECTION.NEW_PAGE)
doc.add_heading("Entity and authority worksheet", level=1)
add_table(
    doc,
    ["Required fact", "Verified value", "Evidence reference"],
    [
        ("Legal entity name", "UNRESOLVED", ""),
        ("Entity type", "UNRESOLVED", ""),
        ("SOS file number", "UNRESOLVED", ""),
        ("Comptroller reference if applicable", "UNRESOLVED", ""),
        ("Requester and authority", "UNRESOLVED", ""),
        ("Authorized signer role", "UNRESOLVED", ""),
        ("Displayed portal organization", "UNRESOLVED", ""),
    ],
    [2.4, 2.6, 2.2],
)

doc.add_heading("Address correction audit", level=1)
address_intro = doc.add_paragraph(
    "Complete one block for every address field. Registered office, principal office, mailing address, "
    "governing-person address, and Comptroller report addresses are separate fields unless current official "
    "instructions expressly say otherwise."
)
address_intro.paragraph_format.space_after = Pt(10)

add_fill_line(doc, "Agency and record")
add_fill_line(doc, "Address role")
add_table(
    doc,
    ["Component", "Current value exactly as shown", "Proposed value", "Evidence ID"],
    [
        ("Street line 1", "UNRESOLVED", "UNRESOLVED", ""),
        ("Street line 2", "UNRESOLVED", "UNRESOLVED", ""),
        ("City", "UNRESOLVED", "UNRESOLVED", ""),
        ("State", "UNRESOLVED", "UNRESOLVED", ""),
        ("ZIP code", "UNRESOLVED", "UNRESOLVED", ""),
        ("County if requested", "UNRESOLVED", "UNRESOLVED", ""),
    ],
    [1.35, 2.55, 2.25, 1.05],
)
add_fill_line(doc, "Reason for correction")
add_fill_line(doc, "Public record or privacy review")

doc.add_heading("Field change log", level=1)
add_table(
    doc,
    ["ID", "Record and field", "Current", "Proposed", "Evidence", "Result"],
    [
        ("CHG-001", "UNRESOLVED", "UNRESOLVED", "UNRESOLVED", "UNRESOLVED", "Blocked"),
        ("CHG-002", "", "", "", "", ""),
        ("CHG-003", "", "", "", "", ""),
    ],
    [0.65, 1.6, 1.35, 1.35, 1.2, 1.05],
)

doc.add_heading("Form routing decision", level=1)
add_table(
    doc,
    ["Facts", "Candidate route", "Decision"],
    [
        ("Unsubmitted portal draft with a data-entry error", "Edit draft and preview", "UNRESOLVED"),
        ("Rejected portal submission", "Edit rejected submission and map rejection reasons", "UNRESOLVED"),
        ("Accepted instrument was inaccurate when filed", "Form 403 may apply", "UNRESOLVED"),
        ("Only later change is registered agent or registered office", "Form 401 may apply", "UNRESOLVED"),
        ("Later amendment to certificate of formation", "Form 424 may apply", "UNRESOLVED"),
        ("PIR, OIR, franchise tax, taxpayer account, or veteran status", "Separate Comptroller handoff", "UNRESOLVED"),
    ],
    [3.2, 2.6, 1.4],
)
add_fill_line(doc, "Selected route")
add_fill_line(doc, "Official instruction URL")
add_fill_line(doc, "Official source checked at")
add_fill_line(doc, "Applicability reviewed by")
add_fill_line(doc, "Fee and waiver result")

doc.add_heading("Route-specific checks", level=2)
for item in (
    "Form 401: current agent and office verified; agent consent verified; Texas service address validated",
    "Form 403: original instrument and filing date identified; correction scope supported by original facts",
    "Form 424: entity-specific approval and signer requirements verified; added, altered, or deleted provisions stated exactly",
    "Comptroller: report type and year verified; SOS and Comptroller effects kept separate",
):
    add_checkbox_line(doc, item)

doc.add_page_break()
doc.add_heading("Evidence map", level=1)
add_table(
    doc,
    ["Evidence ID", "Supports", "Verification", "Restricted", "Reviewer result"],
    [
        ("EVD-001", "UNRESOLVED", "UNRESOLVED", "UNRESOLVED", "Pending"),
        ("EVD-002", "", "", "", ""),
        ("EVD-003", "", "", "", ""),
    ],
    [1.0, 1.3, 2.5, 1.1, 1.3],
)

doc.add_heading("Draft quality review", level=1)
for item in (
    "The original source remains unchanged and its digest is recorded",
    "Every current value was transcribed exactly from the source",
    "Every proposed value has an evidence reference",
    "The form revision and official instructions were checked on the preparation date",
    "The editable draft and clean unsigned copy match",
    "No credentials, sensitive identifiers, veteran records, payment data, or unnecessary PII appear",
    "No signature, certification, or attestation was applied by automation",
    "SOS and Comptroller actions are documented separately",
    "All unresolved material fields block readiness",
):
    add_checkbox_line(doc, item)

doc.add_heading("Approval and human submission gate", level=1)
add_table(
    doc,
    ["Gate", "Required result"],
    [
        ("Exact packet version and SHA-256", "UNRESOLVED"),
        ("Internal reviewer decision", "NOT GRANTED"),
        ("Level C filing approval", "NOT GRANTED"),
        ("Level D financial approval if payment is required", "NOT GRANTED"),
        ("Authorized human signer", "UNRESOLVED"),
        ("Authorized human submitter", "UNRESOLVED"),
        ("Human submission", "NOT PERFORMED"),
        ("Agency receipt", "NOT RECORDED"),
        ("Agency disposition", "NOT RECORDED"),
    ],
    [3.9, 3.3],
)

closing = doc.add_paragraph()
closing.add_run("Stop condition. ").bold = True
closing.add_run(
    "If any material identity, source, authority, address, evidence, routing, signature, fee, or approval field "
    "is unresolved, keep the case blocked. READY FOR HUMAN SUBMISSION is not proof of submission or acceptance."
)

doc.add_heading("Official reference pages", level=1)
for url in (
    "https://www.sos.state.tx.us/corp/forms_boc.shtml",
    "https://www.sos.state.tx.us/corp/bf-guide.shtml",
    "https://www.sos.state.tx.us/corp/instructions/401.shtml",
    "https://www.sos.state.tx.us/corp/instructions/403.shtml",
    "https://www.sos.state.tx.us/corp/instructions/424.shtml",
    "https://comptroller.texas.gov/taxes/franchise/veteran-business.php",
    "https://comptroller.texas.gov/taxes/franchise/pir-oir-filing-req.php",
):
    doc.add_paragraph(url, style="List Bullet")

doc.add_heading("Reviewer notes", level=1)
notes = doc.add_table(rows=5, cols=1)
notes.alignment = WD_TABLE_ALIGNMENT.CENTER
notes.autofit = False
set_table_borders(notes)
mark_row_as_header(notes.rows[0])
for row_index, row in enumerate(notes.rows):
    cell = row.cells[0]
    cell.width = Inches(7.2)
    set_cell_margins(cell, top=170, bottom=170)
    cell.text = "Notes" if row_index == 0 else ""
    if row_index == 0:
        set_cell_shading(cell, "1F4E78")
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
            run.font.size = Pt(9)
    else:
        cell.paragraphs[0].add_run("\n\n")

footer = section.footer
footer_p = footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer_p.add_run("Unsigned preparation workbook | Human review and submission required")
footer_run.font.size = Pt(8)
footer_run.font.color.rgb = RGBColor(89, 89, 89)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
